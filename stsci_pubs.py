#!/usr/bin/env python3
"""Filter an STScI weekly publications email down to papers with authors on a roster.

Usage:
    python3 stsci_pubs.py roster.xlsx email.txt
    pbpaste | python3 stsci_pubs.py roster.xlsx        # macOS: copy email body first
"""
import re
import sys
import unicodedata

from openpyxl import load_workbook


def key(name):
    """'Koekemoer, A. M.' and 'Koekemoer, Anton M' both -> ('koekemoer', 'a')."""
    name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower()
    last, _, first = name.partition(",")
    return last.strip(), first.strip()[:1]


def load_roster(path):
    ws = load_workbook(path, read_only=True).active
    names = [row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
    return {key(n): n for n in names}


def parse_papers(text):
    for block in re.split(r"^\s*-\s*$", text, flags=re.M):
        lines = [l.strip() for l in block.splitlines() if l.strip()]
        m = re.search(r"\(STScI authors: (.*)\)\s*$", block, flags=re.M)
        url_i = next((i for i, l in enumerate(lines) if l.startswith("http")), None)
        if not m or url_i is None:
            continue  # skips the footer and anything malformed
        yield {
            "title": " ".join(lines[1:url_i]),  # tolerates wrapped titles
            "url": lines[url_i],
            "first": lines[url_i + 1].split(",")[0],
            "stsci": [a.strip() for a in m.group(1).split(";")],
        }


def main():
    roster = load_roster(sys.argv[1])
    text = open(sys.argv[2], encoding="utf-8").read() if len(sys.argv) > 2 else sys.stdin.read()

    hits = []
    for p in parse_papers(text):
        matched = [roster[key(a)] for a in p["stsci"] if key(a) in roster]
        if matched:
            hits.append((p, matched))

    if not hits:
        print("No papers this week with authors on the list.")
        return
    print(f"{len(hits)} paper(s) this week with authors on the list:\n")
    for p, matched in hits:
        print(f"{p['title']}\n  {p['first']} et al.\n  {p['url']}\n  HST Research Staff: {'; '.join(matched)}\n")


if __name__ == "__main__":
    main()